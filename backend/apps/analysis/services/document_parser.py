"""
Document Parser — Extracts text and metadata from PDF, DOCX, PPTX files.
Uses PyMuPDF for PDF, python-docx for DOCX, python-pptx for PPTX.
"""
import os
import logging
from dataclasses import dataclass, field
from typing import Optional

logger = logging.getLogger(__name__)


@dataclass
class PageInfo:
    """Information about a single page/slide."""
    page_number: int
    text: str = ''
    has_images: bool = False
    has_tables: bool = False
    image_count: int = 0
    fonts: list = field(default_factory=list)
    width: float = 0
    height: float = 0
    is_image_only: bool = False
    warnings: list = field(default_factory=list)


@dataclass
class DocumentData:
    """Structured data extracted from a document."""
    text: str = ''
    pages: list = field(default_factory=list)  # List[PageInfo]
    page_count: int = 0
    word_count: int = 0
    fonts_used: list = field(default_factory=list)
    has_images: bool = False
    has_tables: bool = False
    has_hyperlinks: bool = False
    headings: list = field(default_factory=list)
    metadata: dict = field(default_factory=dict)
    errors: list = field(default_factory=list)
    warnings: list = field(default_factory=list)


def parse_document(file_path: str, file_type: str) -> DocumentData:
    """Parse a document and extract text and metadata."""
    file_type = file_type.upper()

    if file_type == 'PDF':
        return parse_pdf(file_path)
    elif file_type == 'DOCX':
        return parse_docx(file_path)
    elif file_type in ('PPTX', 'PPT'):
        return parse_pptx(file_path)
    else:
        raise ValueError(f"Unsupported file type: {file_type}")


def parse_pdf(file_path: str) -> DocumentData:
    """Parse PDF using PyMuPDF (fitz)."""
    import fitz  # PyMuPDF

    doc_data = DocumentData()
    all_text_parts = []
    all_fonts = set()

    try:
        doc = fitz.open(file_path)
        doc_data.page_count = len(doc)
        doc_data.metadata = dict(doc.metadata) if doc.metadata else {}

        for page_num in range(len(doc)):
            page = doc[page_num]
            page_info = PageInfo(page_number=page_num + 1)

            # Get page dimensions
            page_info.width = page.rect.width
            page_info.height = page.rect.height

            # Extract text
            text = page.get_text("text")
            page_info.text = text.strip()
            all_text_parts.append(page_info.text)

            # Check for images
            images = page.get_images(full=True)
            page_info.image_count = len(images)
            page_info.has_images = len(images) > 0
            if page_info.has_images:
                doc_data.has_images = True

            # Check if page is image-only (has images but no/little text)
            if page_info.has_images and len(page_info.text.strip()) < 20:
                page_info.is_image_only = True
                page_info.warnings.append("This page appears to be image-only. Text may not be extractable by automated parsers.")

            # Extract font information
            try:
                blocks = page.get_text("dict")["blocks"]
                for block in blocks:
                    if "lines" in block:
                        for line in block["lines"]:
                            for span in line["spans"]:
                                font_name = span.get("font", "")
                                if font_name:
                                    all_fonts.add(font_name)
                                    if font_name not in page_info.fonts:
                                        page_info.fonts.append(font_name)
            except Exception as e:
                logger.warning(f"Font extraction error on page {page_num + 1}: {e}")

            # Check for links
            links = page.get_links()
            if links:
                doc_data.has_hyperlinks = True

            # Check for tables (heuristic: look for structured text blocks)
            tables = page.find_tables()
            if tables and len(tables.tables) > 0:
                page_info.has_tables = True
                doc_data.has_tables = True

            doc_data.pages.append(page_info)

        doc.close()

    except Exception as e:
        logger.error(f"PDF parsing error: {e}", exc_info=True)
        doc_data.errors.append(f"Error parsing PDF: {str(e)}")

    doc_data.text = '\n\n'.join(all_text_parts)
    doc_data.word_count = len(doc_data.text.split())
    doc_data.fonts_used = list(all_fonts)

    # OCR fallback for image-only documents
    if doc_data.word_count < 10 and doc_data.has_images:
        ocr_text = try_ocr_pdf(file_path)
        if ocr_text:
            doc_data.text = ocr_text
            doc_data.word_count = len(ocr_text.split())
            doc_data.warnings.append("Text was extracted using OCR. Results may be less accurate.")

    return doc_data


def try_ocr_pdf(file_path: str) -> Optional[str]:
    """Attempt OCR on a PDF using Tesseract."""
    try:
        import fitz
        import pytesseract
        from PIL import Image
        import io

        doc = fitz.open(file_path)
        ocr_texts = []

        for page_num in range(min(len(doc), 10)):  # Limit to first 10 pages
            page = doc[page_num]
            pix = page.get_pixmap(matrix=fitz.Matrix(2, 2))  # 2x zoom for better OCR
            img = Image.open(io.BytesIO(pix.tobytes("png")))
            text = pytesseract.image_to_string(img)
            if text.strip():
                ocr_texts.append(text.strip())

        doc.close()
        return '\n\n'.join(ocr_texts) if ocr_texts else None

    except ImportError:
        logger.info("Tesseract/pytesseract not available for OCR fallback.")
        return None
    except Exception as e:
        logger.warning(f"OCR fallback failed: {e}")
        return None


def parse_docx(file_path: str) -> DocumentData:
    """Parse DOCX using python-docx."""
    from docx import Document
    from docx.enum.text import WD_ALIGN_PARAGRAPH

    doc_data = DocumentData()
    all_text_parts = []
    all_fonts = set()

    try:
        doc = Document(file_path)

        # Extract core properties
        try:
            core = doc.core_properties
            doc_data.metadata = {
                'title': core.title or '',
                'author': core.author or '',
                'created': str(core.created) if core.created else '',
            }
        except Exception:
            pass

        # Process paragraphs
        for para in doc.paragraphs:
            text = para.text.strip()
            if text:
                all_text_parts.append(text)

            # Detect headings
            if para.style and para.style.name and 'Heading' in para.style.name:
                doc_data.headings.append({
                    'level': para.style.name,
                    'text': text,
                })

            # Collect fonts
            for run in para.runs:
                if run.font and run.font.name:
                    all_fonts.add(run.font.name)

        # Process tables
        for table in doc.tables:
            doc_data.has_tables = True
            for row in table.rows:
                row_text = ' | '.join(cell.text.strip() for cell in row.cells if cell.text.strip())
                if row_text:
                    all_text_parts.append(row_text)

        # Check for images
        for rel in doc.part.rels.values():
            if "image" in rel.reltype:
                doc_data.has_images = True
                break

        # Check for hyperlinks
        for para in doc.paragraphs:
            for run in para.runs:
                if run.font and run.font.color and run.font.underline:
                    # Heuristic: underlined colored text is often a hyperlink
                    pass
            # Check paragraph XML for hyperlinks
            if '<a:hlinkClick' in para._element.xml or '<w:hyperlink' in para._element.xml:
                doc_data.has_hyperlinks = True

        # Sections / headers / footers
        for section in doc.sections:
            try:
                header = section.header
                if header and header.paragraphs:
                    header_text = ' '.join(p.text for p in header.paragraphs if p.text.strip())
                    if header_text:
                        doc_data.warnings.append(f"Header content detected: '{header_text[:50]}...'")
            except Exception:
                pass

            try:
                footer = section.footer
                if footer and footer.paragraphs:
                    footer_text = ' '.join(p.text for p in footer.paragraphs if p.text.strip())
                    if footer_text:
                        doc_data.warnings.append(f"Footer content detected: '{footer_text[:50]}...'")
            except Exception:
                pass

        doc_data.page_count = max(1, len(doc_data.text.split('\n')) // 40)  # Estimate

    except Exception as e:
        logger.error(f"DOCX parsing error: {e}", exc_info=True)
        doc_data.errors.append(f"Error parsing DOCX: {str(e)}")

    doc_data.text = '\n'.join(all_text_parts)
    doc_data.word_count = len(doc_data.text.split())
    doc_data.fonts_used = list(all_fonts)

    return doc_data


def parse_pptx(file_path: str) -> DocumentData:
    """Parse PPTX using python-pptx."""
    from pptx import Presentation
    from pptx.util import Pt, Emu

    doc_data = DocumentData()
    all_text_parts = []
    all_fonts = set()

    try:
        prs = Presentation(file_path)
        doc_data.page_count = len(prs.slides)

        for slide_num, slide in enumerate(prs.slides, 1):
            page_info = PageInfo(page_number=slide_num)
            slide_texts = []

            for shape in slide.shapes:
                # Text frames
                if shape.has_text_frame:
                    for paragraph in shape.text_frame.paragraphs:
                        para_text = paragraph.text.strip()
                        if para_text:
                            slide_texts.append(para_text)

                        # Font analysis
                        for run in paragraph.runs:
                            if run.font:
                                if run.font.name:
                                    all_fonts.add(run.font.name)
                                    if run.font.name not in page_info.fonts:
                                        page_info.fonts.append(run.font.name)

                                # Check for very small text
                                if run.font.size and run.font.size < Pt(8):
                                    page_info.warnings.append(
                                        f"Very small text ({run.font.size.pt}pt): '{para_text[:30]}...'"
                                    )

                # Tables
                if shape.has_table:
                    page_info.has_tables = True
                    doc_data.has_tables = True
                    for row in shape.table.rows:
                        row_text = ' | '.join(
                            cell.text.strip() for cell in row.cells if cell.text.strip()
                        )
                        if row_text:
                            slide_texts.append(row_text)

                # Images
                if shape.shape_type and str(shape.shape_type) in ('13', 'PICTURE', '17'):
                    page_info.has_images = True
                    doc_data.has_images = True
                elif hasattr(shape, 'image'):
                    page_info.has_images = True
                    doc_data.has_images = True

            page_info.text = '\n'.join(slide_texts)

            # Detect image-only slides
            if page_info.has_images and len(page_info.text.strip()) < 10:
                page_info.is_image_only = True
                page_info.warnings.append(
                    "This slide appears to contain mostly images. Important content inside images may not be extractable."
                )

            all_text_parts.append(f"--- Slide {slide_num} ---\n{page_info.text}")
            doc_data.pages.append(page_info)

    except Exception as e:
        logger.error(f"PPTX parsing error: {e}", exc_info=True)
        doc_data.errors.append(f"Error parsing PPTX: {str(e)}")

    doc_data.text = '\n\n'.join(all_text_parts)
    doc_data.word_count = len(doc_data.text.split())
    doc_data.fonts_used = list(all_fonts)

    return doc_data
